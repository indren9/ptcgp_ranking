from __future__ import annotations

import csv
from dataclasses import asdict, dataclass
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import shutil
import tempfile
from typing import Any, Callable, Iterable, Mapping

from openpyxl import load_workbook

from scripts.latest_completed_meta import PUBLIC_RANKING_COLUMNS, validate_bundle
from storage.routing import base_for_expansion


ONEDRIVE_META_ROOT_ENV = "PTCGP_ONEDRIVE_META_ROOT"
GAME_ALIASES = {
    "POCKET": "POCKET",
    "PTCG": "TCG",
}

POCKET_REQUIRED_PATHS = (
    "rankings/mars/mars_ranking_latest.csv",
    "matrices/heatmaps/wr_heatmap_latest.png",
    "reports/mars/mars_matchup_report_latest.xlsx",
    "run/run_manifest_latest.json",
)
POCKET_WILDCARD_PATH = "diagnostics/wildcards/wildcard_candidates_latest.csv"
POCKET_ALLOWED_PATHS = frozenset((*POCKET_REQUIRED_PATHS, POCKET_WILDCARD_PATH))
TCG_PUBLIC_PATHS = (
    "ranking.csv",
    "heatmap.png",
    "manifest.json",
)


class PublicationError(RuntimeError):
    """Raised when publication cannot complete safely."""


@dataclass(frozen=True)
class PublicationArtifact:
    source: Path
    relative_path: str
    kind: str
    manifest: bool = False


@dataclass(frozen=True)
class PublicationResult:
    game: str
    canonical_game: str
    target_root: str
    status: str
    published: tuple[str, ...]
    skipped_identical: tuple[str, ...]
    source_hashes: Mapping[str, str]
    events: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


FailureInjector = Callable[[str, PublicationArtifact], None]


def onedrive_publication_enabled(cfg: Mapping[str, Any] | None) -> bool:
    publication = ((cfg or {}).get("publication") or {})
    onedrive = publication.get("onedrive") or {}
    return bool(onedrive.get("enabled", False))


def canonical_game_root(
    game: str,
    *,
    meta_root: Path | str | None = None,
) -> Path:
    game_key = str(game or "").strip().upper()
    try:
        alias = GAME_ALIASES[game_key]
    except KeyError as exc:
        raise PublicationError(f"Unsupported publication game: {game!r}") from exc

    configured = meta_root or os.environ.get(ONEDRIVE_META_ROOT_ENV)
    root = (
        Path(configured).expanduser()
        if configured
        else Path.home() / "OneDrive" / "Giochi" / "Pokemon" / "PTCGP" / "Meta"
    )
    return (root / alias).resolve()


def _relative_path(value: str) -> str:
    normalized = str(value).replace("\\", "/")
    path = PurePosixPath(normalized)
    if path.is_absolute() or not path.parts or any(part in {"", ".", ".."} for part in path.parts):
        raise PublicationError(f"Unsafe publication relative path: {value!r}")
    return path.as_posix()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_csv(path: Path, *, exact_public_columns: bool = True) -> None:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            columns = tuple(reader.fieldnames or ())
            rows = list(reader)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise PublicationError(f"Ranking CSV is not readable: {path}") from exc

    if exact_public_columns:
        if columns != PUBLIC_RANKING_COLUMNS:
            raise PublicationError(f"Ranking CSV columns do not match the public contract: {path}")
    elif not {"Rank", "Deck"}.issubset(columns):
        raise PublicationError(f"CSV lacks expected core columns: {path}")
    if not rows:
        raise PublicationError(f"CSV has no data rows: {path}")


def _csv_has_data(path: Path) -> bool:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle)
            next(reader, None)
            return next(reader, None) is not None
    except (OSError, UnicodeError, csv.Error) as exc:
        raise PublicationError(f"Wildcard CSV is not readable: {path}") from exc


def _validate_png(path: Path) -> None:
    if path.stat().st_size <= 8 or path.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n":
        raise PublicationError(f"PNG is empty or has an invalid signature: {path}")


def _validate_xlsx(path: Path) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        if not workbook.sheetnames:
            raise PublicationError(f"XLSX contains no worksheets: {path}")
        workbook.close()
    except PublicationError:
        raise
    except Exception as exc:
        raise PublicationError(f"XLSX is not a readable workbook: {path}") from exc


def _load_manifest(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise PublicationError(f"Manifest is not valid JSON: {path}") from exc
    if not isinstance(payload, dict):
        raise PublicationError(f"Manifest root must be an object: {path}")
    return payload


def _validate_recorded_hashes(
    manifest: Mapping[str, Any],
    artifacts: Iterable[PublicationArtifact],
) -> None:
    by_name = {artifact.source.name: artifact for artifact in artifacts if not artifact.manifest}
    outputs = manifest.get("outputs") or {}
    if not isinstance(outputs, Mapping):
        return
    for metadata in outputs.values():
        if not isinstance(metadata, Mapping) or not metadata.get("sha256"):
            continue
        filename = str(
            metadata.get("path")
            or metadata.get("filename")
            or metadata.get("file")
            or ""
        ).replace("\\", "/").split("/")[-1]
        artifact = by_name.get(filename)
        if artifact is None:
            continue
        expected = str(metadata["sha256"]).lower()
        if sha256_file(artifact.source) != expected:
            raise PublicationError(f"Manifest hash mismatch for {filename}")


def _validate_artifact(artifact: PublicationArtifact) -> dict[str, Any] | None:
    path = artifact.source
    if not path.is_file() or path.stat().st_size == 0:
        raise PublicationError(f"Source artifact is missing or empty: {path}")
    if artifact.kind == "ranking_csv":
        _validate_csv(path)
    elif artifact.kind == "png":
        _validate_png(path)
    elif artifact.kind == "xlsx":
        _validate_xlsx(path)
    elif artifact.kind == "manifest":
        return _load_manifest(path)
    elif artifact.kind == "wildcard_csv":
        _validate_csv(path, exact_public_columns=False)
    else:
        raise PublicationError(f"Unsupported artifact kind: {artifact.kind!r}")
    return None


def _copy_source_to_temp(source: Path, destination: Path) -> Path:
    fd, temp_name = tempfile.mkstemp(
        dir=destination.parent,
        prefix=f".{destination.name}.onedrive-publish-",
        suffix=".tmp",
    )
    temp = Path(temp_name)
    try:
        with source.open("rb") as source_handle, os.fdopen(fd, "wb") as temp_handle:
            shutil.copyfileobj(source_handle, temp_handle, length=1024 * 1024)
            temp_handle.flush()
            os.fsync(temp_handle.fileno())
    except Exception:
        temp.unlink(missing_ok=True)
        raise
    return temp


def publish_transaction(
    *,
    game: str,
    source_root: Path,
    target_root: Path,
    artifacts: Iterable[PublicationArtifact],
    allowed_paths: Iterable[str],
    manifest_relative_path: str,
    failure_injector: FailureInjector | None = None,
) -> PublicationResult:
    source_root = source_root.resolve()
    target_root = target_root.resolve()
    allowed = {_relative_path(path) for path in allowed_paths}
    manifest_relative_path = _relative_path(manifest_relative_path)
    items = tuple(artifacts)

    if not source_root.is_dir():
        raise PublicationError(f"Source root is not a directory: {source_root}")
    if not items:
        raise PublicationError("Publication artifact set is empty")

    seen: set[str] = set()
    manifests: list[PublicationArtifact] = []
    manifest_payload: dict[str, Any] | None = None
    normalized_items: list[PublicationArtifact] = []
    source_hashes: dict[str, str] = {}

    for item in items:
        relative = _relative_path(item.relative_path)
        if relative not in allowed:
            raise PublicationError(f"Unexpected artifact rejected: {relative}")
        if relative in seen:
            raise PublicationError(f"Duplicate publication path: {relative}")
        seen.add(relative)
        resolved_source = item.source.resolve()
        try:
            resolved_source.relative_to(source_root)
        except ValueError as exc:
            raise PublicationError(f"Artifact escapes source root: {resolved_source}") from exc
        normalized = PublicationArtifact(resolved_source, relative, item.kind, item.manifest)
        payload = _validate_artifact(normalized)
        if normalized.manifest:
            manifests.append(normalized)
            manifest_payload = payload
        normalized_items.append(normalized)
        source_hashes[relative] = sha256_file(resolved_source)

    if len(manifests) != 1 or manifests[0].relative_path != manifest_relative_path:
        raise PublicationError("Exactly one declared manifest must match manifest_relative_path")
    if manifest_payload is not None:
        _validate_recorded_hashes(manifest_payload, normalized_items)

    ordered = [item for item in normalized_items if not item.manifest] + manifests
    destination_state: dict[str, str | None] = {}
    staged: dict[str, Path] = {}
    published: list[str] = []
    skipped: list[str] = []
    events: list[str] = []

    try:
        for item in ordered:
            destination = target_root / Path(*PurePosixPath(item.relative_path).parts)
            current = sha256_file(destination) if destination.is_file() else None
            destination_state[item.relative_path] = current
            if current == source_hashes[item.relative_path]:
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            temp = _copy_source_to_temp(item.source, destination)
            staged[item.relative_path] = temp
            if failure_injector is not None:
                failure_injector("after_temp_copy", item)
            if sha256_file(temp) != source_hashes[item.relative_path]:
                raise PublicationError(f"Source/temp SHA-256 mismatch: {item.relative_path}")
            events.append(f"TEMP_VERIFIED:{item.relative_path}")

        for item in ordered:
            relative = item.relative_path
            destination = target_root / Path(*PurePosixPath(relative).parts)
            if relative not in staged:
                skipped.append(relative)
                events.append(f"SKIP_IDENTICAL:{relative}")
                continue
            expected_previous = destination_state[relative]
            actual_previous = sha256_file(destination) if destination.is_file() else None
            if actual_previous != expected_previous:
                raise PublicationError(f"Destination changed during publication: {relative}")
            os.replace(staged.pop(relative), destination)
            if sha256_file(destination) != source_hashes[relative]:
                raise PublicationError(f"Canonical SHA-256 mismatch after replace: {relative}")
            published.append(relative)
            events.append(f"CANONICAL_VERIFIED:{relative}")
            if failure_injector is not None:
                failure_injector("after_canonical_replace", item)
    except Exception as exc:
        for temp in staged.values():
            temp.unlink(missing_ok=True)
        if isinstance(exc, PublicationError):
            raise
        raise PublicationError(f"Publication transaction failed: {exc}") from exc

    canonical = GAME_ALIASES[str(game).strip().upper()]
    return PublicationResult(
        game=str(game).strip().upper(),
        canonical_game=canonical,
        target_root=str(target_root),
        status="PASS",
        published=tuple(published),
        skipped_identical=tuple(skipped),
        source_hashes=source_hashes,
        events=tuple(events),
    )


def _pocket_manifest_identity(
    manifest: Mapping[str, Any],
    *,
    format_code: str,
    set_code: str,
    set_name: str,
) -> None:
    source_scope = tuple(str(value) for value in (manifest.get("source_scope") or ()))
    source = ((manifest.get("config_summary") or {}).get("source") or {})
    manifest_set = manifest.get("set") or {}
    if str(source.get("game") or "").upper() != "POCKET":
        raise PublicationError("Pocket manifest game mismatch")
    if len(source_scope) < 2 or source_scope[0].upper() != "POCKET" or source_scope[1].lower() != format_code.lower():
        raise PublicationError("Pocket manifest format/scope mismatch")
    if str(manifest_set.get("code") or "").casefold() != set_code.casefold():
        raise PublicationError("Pocket manifest set code mismatch")
    if str(manifest_set.get("name") or "").strip().casefold() != set_name.strip().casefold():
        raise PublicationError("Pocket manifest set name mismatch")


def publish_pocket_result(
    result: Any,
    *,
    meta_root: Path | str | None = None,
    failure_injector: FailureInjector | None = None,
) -> PublicationResult:
    source_run = base_for_expansion(result.paths.outputs, result.expansion)
    scope = tuple(source_run.relative_to(result.paths.output_root).parts)
    if len(scope) != 3 or scope[0].upper() != "POCKET":
        raise PublicationError(f"Unexpected Pocket output scope: {scope!r}")
    format_code, set_directory = scope[1], scope[2]
    set_code = str(getattr(result.expansion, "code", "") or "").strip()
    set_name = str(getattr(result.expansion, "name", "") or "").strip()
    if not set_code or not set_name or not set_directory.startswith(f"{set_code}__"):
        raise PublicationError("Pocket expansion identity is incomplete or inconsistent")

    specs = (
        ("mars_ranking", POCKET_REQUIRED_PATHS[0], "ranking_csv", False),
        ("heatmap_topN_latest", POCKET_REQUIRED_PATHS[1], "png", False),
        ("report_latest", POCKET_REQUIRED_PATHS[2], "xlsx", False),
        ("run_manifest", POCKET_REQUIRED_PATHS[3], "manifest", True),
    )
    artifacts: list[PublicationArtifact] = []
    for key, relative, kind, manifest in specs:
        source = result.outputs.get(key)
        if source is None:
            raise PublicationError(f"Pocket run lacks required final artifact: {key}")
        artifacts.append(PublicationArtifact(Path(source), relative, kind, manifest))

    wildcard = result.outputs.get("wildcard_candidates")
    if wildcard is not None and Path(wildcard).is_file() and _csv_has_data(Path(wildcard)):
        artifacts.append(PublicationArtifact(Path(wildcard), POCKET_WILDCARD_PATH, "wildcard_csv"))

    manifest = _load_manifest(Path(result.outputs["run_manifest"]))
    _pocket_manifest_identity(
        manifest,
        format_code=format_code,
        set_code=set_code,
        set_name=set_name,
    )
    target = canonical_game_root("POCKET", meta_root=meta_root) / format_code / set_directory
    return publish_transaction(
        game="POCKET",
        source_root=source_run,
        target_root=target,
        artifacts=artifacts,
        allowed_paths=POCKET_ALLOWED_PATHS,
        manifest_relative_path=POCKET_REQUIRED_PATHS[3],
        failure_injector=failure_injector,
    )


def publish_tcg_bundle(
    *,
    bundle_dir: Path,
    plan: Mapping[str, Any],
    meta_root: Path | str | None = None,
    failure_injector: FailureInjector | None = None,
) -> PublicationResult:
    validation_plan = {
        "action": "publish",
        "completed_set": dict(plan.get("latest_completed_window") or {}),
    }
    manifest = validate_bundle(bundle_dir, validation_plan)
    if str(manifest.get("game") or "").upper() != "PTCG":
        raise PublicationError("TCG bundle manifest game mismatch")
    if str(manifest.get("format") or "").lower() != "standard":
        raise PublicationError("TCG bundle manifest format mismatch")

    artifacts = (
        PublicationArtifact(bundle_dir / "ranking.csv", "ranking.csv", "ranking_csv"),
        PublicationArtifact(bundle_dir / "heatmap.png", "heatmap.png", "png"),
        PublicationArtifact(bundle_dir / "manifest.json", "manifest.json", "manifest", True),
    )
    return publish_transaction(
        game="PTCG",
        source_root=bundle_dir,
        target_root=canonical_game_root("PTCG", meta_root=meta_root),
        artifacts=artifacts,
        allowed_paths=TCG_PUBLIC_PATHS,
        manifest_relative_path="manifest.json",
        failure_injector=failure_injector,
    )


__all__ = [
    "GAME_ALIASES",
    "ONEDRIVE_META_ROOT_ENV",
    "POCKET_ALLOWED_PATHS",
    "POCKET_REQUIRED_PATHS",
    "POCKET_WILDCARD_PATH",
    "PublicationArtifact",
    "PublicationError",
    "PublicationResult",
    "TCG_PUBLIC_PATHS",
    "canonical_game_root",
    "onedrive_publication_enabled",
    "publish_pocket_result",
    "publish_tcg_bundle",
    "publish_transaction",
    "sha256_file",
]
