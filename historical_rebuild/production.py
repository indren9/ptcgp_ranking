"""Thin adapters around the frozen Tournament API / Core / MARS implementation."""
from __future__ import annotations

import ast
from copy import deepcopy
from dataclasses import asdict
from importlib.metadata import version
from io import StringIO
import inspect
from pathlib import Path
import platform
import re
import textwrap
import zipfile

import pandas as pd
import requests

from acquisition.scope import EligibilityPolicy, ScopePolicy
from .model import RebuildError, RetryableError, ValidationError, digest, utc
from .store import atomic_write, file_digest

BASE = Path(__file__).resolve().parents[1]


def function_semantics(filename: str, names: list[str]) -> str:
    """Hash ASTs plus recursively referenced module-local functions/constants.

    A report-only edit in the monolithic deck_ranking module must not invalidate
    Core. External dependencies are enumerated separately below.
    """
    tree = ast.parse((BASE / filename).read_text(encoding="utf-8"))
    symbols = {}
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            symbols[node.name] = node
        elif isinstance(node, (ast.Assign, ast.AnnAssign)):
            for target in (node.targets if isinstance(node, ast.Assign) else [node.target]):
                if isinstance(target, ast.Name):
                    symbols[target.id] = node
    pending = list(names)
    selected = {}
    while pending:
        name = pending.pop()
        if name in selected or name not in symbols:
            continue
        node = symbols[name]
        selected[name] = ast.dump(node, include_attributes=False)
        pending.extend(n.id for n in ast.walk(node) if isinstance(n, ast.Name) and n.id in symbols)
    if any(name not in selected for name in names):
        raise RebuildError(f"cannot fingerprint required implementation: {filename}/{names}")
    return digest(selected)


def pack(frame: pd.DataFrame) -> dict:
    # Retain float representations and explicit axes, including numeric-looking
    # canonical IDs whose leading zeroes pandas otherwise infers away.
    return {"csv": frame.to_csv(index=True, lineterminator="\n"),
            "columns": list(frame.columns), "index_name": frame.index.name,
            "index": list(frame.index),
            "string_columns": [c for c in frame.columns if str(frame[c].dtype) in {"object", "string"}]}


def unpack(data: dict) -> pd.DataFrame:
    frame = pd.read_csv(StringIO(data["csv"]), index_col=0, float_precision="round_trip",
                        dtype={c: str for c in data["string_columns"]}, keep_default_na=False,
                        na_values=[""])
    frame.index = pd.Index(data["index"], name=data["index_name"])
    return frame


def stable_workbook(path: Path):
    """Remove container creation times only; leave all workbook data untouched."""
    from io import BytesIO
    output = BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            raw = source.read(name)
            if name == "docProps/core.xml":
                raw = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                             rb"\g<1>2000-01-01T00:00:00Z\g<2>", raw)
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, raw)
    atomic_write(path, output.getvalue())


class ProductionBackend:
    def __init__(self, cfg: dict, *, client=None):
        self.cfg = deepcopy(cfg)
        self.api = self.cfg["source"]["tournament_api"]
        self.eligibility = EligibilityPolicy(game="PTCG", **self.api["eligibility"])
        if self.cfg["source"].get("game") != "PTCG":
            raise RebuildError("historical runner requires the PTCG production configuration")
        self.client = client  # Lazy: STATUS never constructs a requests session.

    def semantics(self, stage):
        methods = {
            "WINDOW_READY": [], "DISCOVERY": ["_discover", "_scope", "_client", "_network"],
            "TOURNAMENT_IDS_FROZEN": ["_freeze", "_scope", "_client", "_network"],
            "RAW_ACQUISITION": ["acquire", "validate_raw", "_scope", "_client", "_network"],
            "NORMALIZATION": ["_normalize"], "CORE": ["_core"], "MARS": ["_mars"],
            "VALIDATION": ["_validate"], "REPORT": ["_report"], "DONE": [],
        }
        own = {name: ast.dump(ast.parse(textwrap.dedent(inspect.getsource(getattr(type(self), name)))))
               for name in methods[stage]}
        files, functions, config, packages = [], {}, {}, []
        if stage in {"DISCOVERY", "TOURNAMENT_IDS_FROZEN"}:
            config = {"eligibility": asdict(self.eligibility)}
            files = ["acquisition/scope.py", "acquisition/selection.py", "domain/releases.py"]
            functions["pipelines/limitless_api_acquisition.py"] = ["_canonical_discovery_records", "_discovery_candidates", "_selection_record"]
            if stage == "DISCOVERY":
                config["discovery"] = {k: self.api.get(k) for k in ("discovery_page_size", "discovery_max_pages")}
        if stage in {"DISCOVERY", "TOURNAMENT_IDS_FROZEN", "RAW_ACQUISITION"}:
            files += ["sources/limitless/tournament_api/client.py", "sources/limitless/tournament_api/release_catalog.py"]
            packages = ["requests"]
        if stage == "NORMALIZATION":
            # Production Tournament API uses canonical IDs and bypasses aliases.
            # Record that policy; do not hash the irrelevant legacy alias file.
            config = {"alias_policy": "canonical_deck_ids_no_legacy_aliases_v1"}
            files = ["sources/limitless/tournament_api/normalize.py", "acquisition/aggregation.py",
                     "acquisition/contracts.py", "acquisition/production_bridge.py",
                     "sources/limitless/tournament_api/release_catalog.py"]
            functions["pipelines/limitless_api_acquisition.py"] = ["_concat_or_empty"]
            packages = ["pandas", "numpy"]
        if stage == "CORE":
            config = {k: self.cfg.get(k) for k in ("top_meta", "analysis", "nan_filter", "wr_policy")}
            files = ["core/" + name for name in ("normalize.py", "consolidate.py", "matrices.py", "nan_filter.py", "nan_diagnostics.py")]
            functions["pipelines/deck_ranking.py"] = ["_build_core_matrices"]
            functions["storage/writers.py"] = ["write_csv_versioned"]
            files += ["storage/routing.py", "storage/paths.py", "domain/expansions.py"]
            packages = ["pandas", "numpy"]
        if stage == "MARS":
            from mars.config import MARSConfig
            config = asdict(MARSConfig(**self.cfg.get("mars", {})))
            files = ["mars/" + name for name in ("config.py", "pipeline.py", "auto_k_cv.py", "posterior.py", "meta.py", "mas_lb.py", "bt.py", "composite.py", "coverage.py", "validate_io.py")]
            packages = ["pandas", "numpy", "scipy"]
            functions["core/matrices.py"] = ["topmeta_post_alias"]
            functions["core/normalize.py"] = ["apply_alias_series"]
        if stage == "VALIDATION":
            files = ["mars/validate_io.py"]
            packages = ["pandas", "numpy"]
        if stage == "REPORT":
            import matplotlib.font_manager as fm
            import zlib
            config = {"format": "production_matchup_workbook_v1", "stable_container_metadata": True}
            config["fonts"] = {
                weight: file_digest(Path(fm.findfont(fm.FontProperties(family="DejaVu Sans", weight=weight))))
                for weight in ("regular", "bold")}
            config["zlib"] = zlib.ZLIB_RUNTIME_VERSION
            files = ["mars/report.py", "mars/meta.py", "mars/config.py", "reporting/excel.py"]
            own["stable_workbook"] = inspect.getsource(stable_workbook)
            packages = ["pandas", "numpy", "openpyxl", "pillow", "matplotlib"]
            functions["storage/writers.py"] = ["write_excel_versioned"]
            functions["core/matrices.py"] = ["topmeta_post_alias"]
            functions["core/normalize.py"] = ["apply_alias_series"]
        if stage in {"NORMALIZATION", "CORE", "MARS", "VALIDATION", "REPORT"}:
            own.update(pack=inspect.getsource(pack), unpack=inspect.getsource(unpack))
        return {"contract": 1, "python": platform.python_version(), "adapter": digest(own), "config": config,
                "files": {f: file_digest(BASE / f) for f in files},
                "functions": {f: function_semantics(f, names) for f, names in functions.items()},
                "packages": {p: version(p) for p in packages}}

    def _client(self):
        if self.client is None:
            from sources.limitless.tournament_api.client import LimitlessTournamentApiClient
            self.client = LimitlessTournamentApiClient(min_request_interval_seconds=self.api.get("min_request_interval_seconds", 4))
        return self.client

    def close(self):
        if self.client is not None:
            self.client.close()

    def _network(self, method, *args, **kwargs):
        try:
            return getattr(self._client(), method)(*args, **kwargs, use_cache=False)
        except (requests.ConnectionError, requests.Timeout) as exc:
            raise RetryableError(str(exc)) from exc
        except requests.HTTPError as exc:
            status = exc.response.status_code if exc.response is not None else None
            if status == 429 or (status is not None and status >= 500):
                raise RetryableError(str(exc)) from exc
            raise RebuildError(f"acquisition blocked HTTP {status}") from exc

    def _scope(self, window):
        return ScopePolicy("ptcg_tcgl_expansion_window_v1", "PTCG", "STANDARD", window.window_id,
                           window.name, utc(window.start_utc), utc(window.effective_end), "reviewed-boundaries-v1")

    def _discover(self, window, inputs, directory):
        from pipelines.limitless_api_acquisition import _canonical_discovery_records, _discovery_candidates
        rows = self._network("list_tournaments", game="PTCG", format="STANDARD",
                             page_size=self.api.get("discovery_page_size", 50),
                             max_pages=self.api.get("discovery_max_pages", 100))
        rows, _ = _canonical_discovery_records(rows)
        ids, diagnostics = _discovery_candidates(rows, scope=self._scope(window), page_size=self.api.get("discovery_page_size", 50))
        return {"candidate_ids": list(ids), "diagnostics": diagnostics}

    def _freeze(self, window, inputs, directory):
        from acquisition.selection import select_tournaments
        from pipelines.limitless_api_acquisition import _selection_record
        details = {tid: self._network("get_tournament_details", tid) for tid in inputs["DISCOVERY"]["candidate_ids"]}
        if any(str(row.get("id")) != tid for tid, row in details.items()):
            raise RebuildError("discovery/details identity mismatch")
        selection = select_tournaments([_selection_record(d) for d in details.values()],
                                      scope=self._scope(window), eligibility=self.eligibility)
        if selection.exclusion_counts.get("invalid_record") or selection.failures:
            raise RebuildError("invalid eligibility evidence; cannot freeze a complete selection")
        if not selection.tournament_ids:
            raise RebuildError("no eligible tournaments; no historical ranking can be validated")
        return {"tournaments": [{"id": tid, "details": details[tid]} for tid in selection.tournament_ids],
                "exclusion_counts": dict(selection.exclusion_counts)}

    def acquire(self, window, item):
        # Details are already frozen, validated eligibility evidence. Acquire the
        # other two endpoints exactly once per missing/stale tournament.
        return {"details": item["details"],
                "standings": self._network("get_tournament_standings", item["id"]),
                "pairings": self._network("get_tournament_pairings", item["id"])}

    def validate_raw(self, tid, payload):
        if set(payload) != {"details", "standings", "pairings"} or not isinstance(payload["details"], dict) or str(payload["details"].get("id")) != tid:
            raise ValidationError("RAW identity/payload contract failed")
        for key in ("standings", "pairings"):
            if not isinstance(payload[key], list) or not all(isinstance(row, dict) for row in payload[key]):
                raise ValidationError(f"RAW {key} is not an array of records")
        if not payload["standings"] or not payload["pairings"]:
            raise ValidationError("RAW missing standings/pairings evidence")
        if any(not row.get("player") for row in payload["standings"]):
            raise ValidationError("RAW participant identity missing")
        if any("phase" not in row or "round" not in row for row in payload["pairings"]):
            raise ValidationError("RAW pairing structure incomplete")

    def _normalize(self, window, inputs, directory):
        from sources.limitless.tournament_api.normalize import normalize_snapshot, PARTICIPANT_COLUMNS, PAIRING_COLUMNS
        from pipelines.limitless_api_acquisition import _concat_or_empty
        from acquisition.aggregation import aggregate_meta, aggregate_matchups
        from acquisition.contracts import adapt_top_meta_decklist, adapt_matchup_raw, materialize_dense_score, AcquisitionFrames
        from acquisition.production_bridge import bridge_tournament_api_frames
        participants, pairings = [], []
        for item in inputs["RAW_ACQUISITION"]["tournaments"]:
            raw = inputs["load_raw"](item["id"])
            _, players, pairs = normalize_snapshot(tournament_id=item["id"], raw_snapshot_id=item["sha256"], **raw)
            participants.append(players)
            pairings.append(pairs)
        players = _concat_or_empty(participants, PARTICIPANT_COLUMNS)
        pairs = _concat_or_empty(pairings, PAIRING_COLUMNS)
        meta = aggregate_meta(players)
        matches = aggregate_matchups(players, pairs)
        top = adapt_top_meta_decklist(meta.meta)
        match = adapt_matchup_raw(matches.matchups)
        dense = materialize_dense_score(match, tuple(zip(top["Deck ID"], top["Deck"])))
        bridge = bridge_tournament_api_frames(AcquisitionFrames(top, match, dense))
        return {"top": pack(bridge.top_meta_decklist), "dense": pack(bridge.dense_score),
                "identities": pack(bridge.deck_identity_map)}

    def _core(self, window, inputs, directory):
        from pipelines.deck_ranking import _build_core_matrices
        from storage.paths import ProjectPaths
        cfg = deepcopy(self.cfg)
        cfg["saving"] = {"output_profile": "user", "duplicate_legacy_latest": False, "include_time_when_changed": False}
        # Production helper writes only inside this unpublished generation.
        paths = ProjectPaths(directory, directory, directory, directory / "cache", directory / "logs")
        normalized = inputs["NORMALIZATION"]
        frames, _, _ = _build_core_matrices(cfg=cfg, paths=paths, exp=None,
            df_matchup_raw=unpack(normalized["dense"]), df_top_meta=unpack(normalized["top"]),
            preserve_zero_evidence=True, alias_index_override={}, canonical_dense_input=True)
        return {name: pack(frames[name]) for name in ("score_flat", "wr_matrix", "n_dir_matrix")}

    def _mars(self, window, inputs, directory):
        from core.matrices import topmeta_post_alias
        from mars.config import MARSConfig
        from mars.pipeline import run_mars
        core = inputs["CORE"]
        ranking, diagnostics, coverage, missing = run_mars(
            unpack(core["wr_matrix"]), unpack(core["n_dir_matrix"]), unpack(core["score_flat"]),
            topmeta_post_alias(unpack(inputs["NORMALIZATION"]["top"]), {}), MARSConfig(**self.cfg.get("mars", {})))
        ranking.to_csv(directory / "ranking.csv", lineterminator="\n")
        # Full diagnostics may include numpy objects; retain the effective K
        # needed by the report in a strict, deterministic JSON contract.
        return {"ranking": pack(ranking), "coverage": pack(coverage), "missing": pack(missing),
                "K_used": float(diagnostics["AUTO_K"]["K_used"])}

    def _validate(self, window, inputs, directory):
        import numpy as np
        from mars.validate_io import validate_contract
        wr = unpack(inputs["CORE"]["wr_matrix"])
        counts = unpack(inputs["CORE"]["n_dir_matrix"])
        ranking = unpack(inputs["MARS"]["ranking"])
        result = validate_contract(wr, counts)
        if not result["ok"] or result["issues"] or len(wr) < 2 or len(ranking) != len(wr):
            raise ValidationError(f"Core/ranking validation failed: {result}")
        if ranking["Deck"].duplicated().any() or set(ranking["Deck"]) != set(wr.index):
            raise ValidationError("ranking axis does not match Core")
        for key in ("Score_%", "MAS_%", "LB_%", "BT_%", "SE_%"):
            if not np.isfinite(ranking[key].to_numpy(dtype=float)).all():
                raise ValidationError(f"non-finite ranking {key}")
        if not ranking["Score_%"].is_monotonic_decreasing:
            raise ValidationError("ranking order is invalid")
        return {"validated": True, "core_hash": digest(inputs["CORE"]), "mars_hash": digest(inputs["MARS"])}

    def _report(self, window, inputs, directory):
        from core.matrices import topmeta_post_alias
        from mars.report import write_mars_matchup_report
        from mars.config import MARSConfig
        from mars.meta import blend_meta
        from openpyxl import load_workbook
        core = inputs["CORE"]
        wr, counts = unpack(core["wr_matrix"]), unpack(core["n_dir_matrix"])
        cfg = MARSConfig(**self.cfg.get("mars", {}))
        weights, info = blend_meta(list(wr.index), counts,
                                  topmeta_post_alias(unpack(inputs["NORMALIZATION"]["top"]), {}), cfg)
        _, path, _ = write_mars_matchup_report(ranking_df=unpack(inputs["MARS"]["ranking"]),
            filtered_wr=wr, n_dir=counts, p_blend=weights, K_used=inputs["MARS"]["K_used"],
            score_flat=unpack(core["score_flat"]), mu=cfg.MU, gamma=info.get("gamma"),
            out_dir=directory, also_versioned=False, keep_legend_image=False)
        stable_workbook(path)
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if "01_Summary" not in book.sheetnames or book["01_Summary"].max_row < 2:
                raise ValidationError("report workbook validation failed")
        finally:
            book.close()
        return {"validated": True, "workbook": path.name, "sha256": file_digest(path)}

    def execute(self, stage, window, inputs, directory):
        methods = {"DISCOVERY": self._discover, "TOURNAMENT_IDS_FROZEN": self._freeze,
                   "NORMALIZATION": self._normalize, "CORE": self._core, "MARS": self._mars,
                   "VALIDATION": self._validate, "REPORT": self._report}
        return methods[stage](window, inputs, directory)
